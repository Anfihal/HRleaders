"""Convert the committed FSP PDF template into an editable PPTX template.

Only the PDF version of the FSP deck exists in the repository (issue #8). Each
PDF page becomes one 16:9 slide that keeps the exact page geometry:

* a full-slide background picture rendered from the page with its text removed
  (raster artwork and vector shapes stay untouched);
* one text box per PDF text block, at the block's original position, with the
  original size, weight and colour of every span;
* an empty, named text region for every filled card drawn as a vector shape
  without text, so content can be placed inside the template's own cards.

Shape names encode their origin (``text:<block>`` / ``region:<index>``) so the
generator can address editable regions without guessing.
"""

from __future__ import annotations

import argparse
import io
import logging
from pathlib import Path
from typing import Any

import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from tools.fsp_presentation import TEMPLATE_FILENAME

PDF_FILENAME = "ЛЦТ_2026_Шаблон презентации ФСП.pdf"
# The PDF was exported with a fallback face; page 5 of the template names the real one.
FONT = "Montserrat"
BULLETS = {"•", "§", "-"}
BACKGROUND_DPI = 192
EMU_PER_PT = 12700
# Page 1 has its "Шаблон презентации" subtitle baked into the artwork. It is the
# only text of the title slide that the template intends to be replaced, so the
# strip is repainted from the rows above and below it and exposed as a region.
TITLE_SUBTITLE = (30, 476, 300, 504)
# The editable title region spans the empty artwork between the event name and the footer.
TITLE_REGION = (30, 420, 600, 510)
LOG = logging.getLogger(__name__)


def hex_color(value: int) -> RGBColor:
    return RGBColor((value >> 16) & 255, (value >> 8) & 255, value & 255)


def page_lines(page: pymupdf.Page) -> list[dict[str, Any]]:
    """Return text blocks as lists of visual lines; bullet glyphs join their line."""
    blocks = []
    for block in page.get_text("dict")["blocks"]:
        if block["type"] != 0:
            continue
        lines: list[dict[str, Any]] = []
        for line in block["lines"]:
            spans = [s for s in line["spans"] if s["text"]]
            if not spans:
                continue
            middle = (line["bbox"][1] + line["bbox"][3]) / 2
            for previous in lines:
                if abs(previous["middle"] - middle) < 4:
                    previous["spans"].extend(spans)
                    break
            else:
                lines.append({"middle": middle, "spans": spans})
        for line in lines:
            line["spans"].sort(key=lambda s: s["bbox"][0])
        lines.sort(key=lambda line: line["middle"])
        if lines:
            blocks.append({"bbox": block["bbox"], "lines": lines})
    return blocks


def background(page: pymupdf.Page) -> bytes:
    """Render the page without text; keep images and vector graphics."""
    for word in page.get_text("words"):
        page.add_redact_annot(pymupdf.Rect(word[:4]), fill=False)
    page.apply_redactions(
        images=pymupdf.PDF_REDACT_IMAGE_NONE,
        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
        text=pymupdf.PDF_REDACT_TEXT_REMOVE,
    )
    pixmap = page.get_pixmap(dpi=BACKGROUND_DPI)
    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    if page.number == 0:
        scale = BACKGROUND_DPI / 72
        x0, y0, x1, y1 = (round(v * scale) for v in TITLE_SUBTITLE)
        # Interpolate every column between the clean rows just above and below.
        pixels = image.load()
        for x in range(x0, x1):
            top, bottom = pixels[x, y0 - 1], pixels[x, y1 + 1]
            for y in range(y0, y1 + 1):
                t = (y - y0 + 1) / (y1 - y0 + 2)
                mixed = zip(top, bottom, strict=True)
                pixels[x, y] = tuple(round(a + (b - a) * t) for a, b in mixed)
    buffer = io.BytesIO()
    image.save(buffer, "JPEG", quality=90, optimize=True)
    return buffer.getvalue()


def add_box(slide: Any, name: str, rect: tuple[float, float, float, float]) -> Any:
    x0, y0, x1, y1 = rect
    box = slide.shapes.add_textbox(
        Emu(round(x0 * EMU_PER_PT)),
        Emu(round(y0 * EMU_PER_PT)),
        Emu(round((x1 - x0) * EMU_PER_PT)),
        Emu(round((y1 - y0) * EMU_PER_PT)),
    )
    box.name = name
    frame = box.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.auto_size = MSO_AUTO_SIZE.NONE
    return box


def write_block(box: Any, block: dict[str, Any]) -> None:
    frame = box.text_frame
    frame.word_wrap = False
    left = block["bbox"][0]
    bulleted = any(line["spans"][0]["text"].strip() in BULLETS for line in block["lines"])
    indent = 0.0
    for number, line in enumerate(block["lines"]):
        paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
        spans = line["spans"]
        bullet = None
        if spans[0]["text"].strip() in BULLETS and len(spans) > 1:
            bullet, spans = spans[0], spans[1:]
            while len(spans) > 1 and not spans[0]["text"].strip():
                spans = spans[1:]
            indent = spans[0]["bbox"][0] - left
        if number + 1 < len(block["lines"]):
            pitch = block["lines"][number + 1]["middle"] - line["middle"]
            paragraph.line_spacing = Pt(max(pitch, spans[0]["size"]))
        else:
            paragraph.line_spacing = Pt(max(spans[0]["size"] * 1.2, 1))
        properties = paragraph._p.get_or_add_pPr()
        if bulleted and indent:
            properties.set("marL", str(round(indent * EMU_PER_PT)))
            hanging = indent if bullet is not None else 0
            properties.set("indent", str(round(-hanging * EMU_PER_PT)))
        if bullet is not None:
            color = properties.makeelement(qn("a:buClr"), {})
            srgb = color.makeelement(qn("a:srgbClr"), {"val": f"{bullet['color']:06X}"})
            color.append(srgb)
            properties.append(color)
            glyph = "•" if bullet["text"].strip() == "§" else bullet["text"].strip()
            properties.append(properties.makeelement(qn("a:buChar"), {"char": glyph}))
        else:
            properties.append(properties.makeelement(qn("a:buNone"), {}))
        text_spans = [s for s in spans if s["text"]]
        for index, span in enumerate(text_spans):
            text = span["text"]
            if index and span["bbox"][0] - text_spans[index - 1]["bbox"][2] > span["size"] * 0.2:
                text = " " + text.lstrip()
            run = paragraph.add_run()
            run.text = text
            run.font.name = FONT
            run.font.size = Pt(round(span["size"] * 2) / 2)
            run.font.bold = "Bold" in span["font"]
            run.font.color.rgb = hex_color(span["color"])


def card_regions(page: pymupdf.Page, blocks: list[dict[str, Any]]) -> list[pymupdf.Rect]:
    """Filled vector cards that hold no whole template text block are the deck's regions.

    Colour chips whose labels belong to a wider text row (slide 5) count as regions too.
    """
    regions: list[pymupdf.Rect] = []
    for drawing in page.get_drawings():
        rect = drawing["rect"]
        if drawing.get("fill") is None or rect.width < 40 or rect.height < 30:
            continue
        if any(rect.contains(pymupdf.Rect(b["bbox"])) for b in blocks):
            continue
        if any(abs(rect.x0 - r.x0) < 3 and abs(rect.y0 - r.y0) < 3 for r in regions):
            continue
        regions.append(rect)
    return regions


def convert(pdf_path: Path, output: Path) -> None:
    document = pymupdf.open(pdf_path)
    width, height = document[0].rect.width, document[0].rect.height
    deck = Presentation()
    deck.slide_width = Emu(round(width * EMU_PER_PT))
    deck.slide_height = Emu(round(height * EMU_PER_PT))
    blank = deck.slide_layouts[6]
    for page in document:
        LOG.debug("page %d", page.number + 1)
        blocks = page_lines(page)
        regions = card_regions(page, blocks)
        slide = deck.slides.add_slide(blank)
        picture = slide.shapes.add_picture(
            io.BytesIO(background(page)), 0, 0, deck.slide_width, deck.slide_height
        )
        picture.name = "background"
        for index, region in enumerate(regions, 1):
            inset = min(14.0, region.height * 0.15)
            box = add_box(
                slide,
                f"region:{index}",
                (region.x0 + inset, region.y0 + inset, region.x1 - inset, region.y1 - inset),
            )
            box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        if page.number == 0:
            add_box(
                slide, "region:subtitle", TITLE_REGION
            ).text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        for index, block in enumerate(blocks, 1):
            x0, y0, x1, y1 = block["bbox"]
            box = add_box(slide, f"text:{index}", (x0, y0, x1 + 4, y1))
            write_block(box, block)
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(output))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf", type=Path, default=Path(PDF_FILENAME))
    parser.add_argument("--output", type=Path, default=Path("templates") / TEMPLATE_FILENAME)
    parser.add_argument("--verbose", action="store_true")
    args = parser.parse_args(argv)
    if args.verbose:
        logging.basicConfig(level=logging.DEBUG, format="%(levelname)s: %(message)s")
    convert(args.pdf, args.output)
    print(f"PPTX template written: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
